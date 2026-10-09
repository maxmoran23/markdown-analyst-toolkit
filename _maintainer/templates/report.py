from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DATA = None  # __DATA__
# __CASE_VALIDATION__
OUTPUT = Path(DATA['suite_id'] + '-report.docx')
doc = Document()
doc.core_properties.author = 'Markdown Analyst Toolkit'
doc.core_properties.title = DATA['title']
doc.core_properties.subject = DATA['notice']
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.75)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string('243641')
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
for name, size in [('Title',24), ('Heading 1',16), ('Heading 2',12)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string('000000' if name=='Title' else '344E5F')
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(14)
    style.paragraph_format.space_after = Pt(7)
for style in [normal, doc.styles['Title'], doc.styles['Heading 1'], doc.styles['Heading 2'], doc.styles['List Bullet']]:
    style.font.name = 'Arial'
    fonts = style.element.get_or_add_rPr().rFonts
    if fonts is not None:
        for key in list(fonts.attrib):
            if 'theme' in key.lower():
                del fonts.attrib[key]
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

def text(value):
    return 'Not provided' if value is None else str(value).lower() if isinstance(value,bool) else '; '.join(map(str,value)) if isinstance(value,list) else str(value)

def formatted(value, fmt):
    if type(value) not in (int,float) or not fmt:
        return text(value)
    if fmt == '0.0%':
        return f'{value:.1%}'
    if fmt == '0.0"%"':
        return f'{value:.1f}%'
    decimal = re.search(r'\.([0#]+)', fmt)
    places = len(decimal.group(1)) if decimal else 0
    number = f'{value:,}' + ('.'+'0'*places if places else '') if isinstance(value,int) else f'{value:,.{places}f}'
    return ('$' if '$' in fmt else '') + number

def table(headers, rows, keep_together=False, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    if widths:
        t.autofit = False
        for column,width in zip(t.columns,widths):column.width=Inches(width)
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text = h
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells=t.add_row().cells
        for i,v in enumerate(values):
            cells[i].text=text(v)
    for row_index,row in enumerate(t.rows):
        if row_index==0 or max(len(cell.text) for cell in row.cells)<=1200:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell_index,cell in enumerate(row.cells):
            if widths:cell.width=Inches(widths[cell_index])
            shading=OxmlElement('w:shd')
            shading.set(qn('w:fill'),'344E5F' if row_index==0 else 'E2E8E9' if row_index%2 else 'F1F2EF')
            cell._tc.get_or_add_tcPr().append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4)
                if keep_together and row_index<len(t.rows)-1:
                    p.paragraph_format.keep_with_next=True
                for run in p.runs:
                    run.font.name='Arial'
                    run.font.size=Pt(9)
                    run.font.color.rgb=RGBColor.from_string('FFFFFF' if row_index==0 else '243641')
                    run.font.bold=row_index==0
    doc.add_paragraph()
    return t

doc.add_paragraph(DATA['title'], 'Title')
doc.add_paragraph(f"Version {DATA['version']} | As of {DATA['as_of']} | Draft for human review")
doc.add_paragraph(DATA['notice'])
doc.add_paragraph(DATA['summary'])
doc.add_paragraph('Decision supported: '+DATA['decision'])
for index, heading in enumerate(DATA['report_sections']):
    doc.add_heading(heading, level=1)
    for paragraph in DATA['report_narrative'][index]:
        doc.add_paragraph(paragraph)
    if index == 0:
        table(['Measure','Value','Calculation or limit'], [[m['label'],formatted(m['value'],m.get('format')),m['method']] for m in DATA['metrics']])
    elif index == 2:
        for f in DATA['findings']:
            doc.add_heading(f["id"]+' '+f['title'], level=2)
            meta=doc.add_paragraph(f"{f['severity']} severity | {f['confidence']} confidence | {f['claim_class']} | Evidence: {', '.join(f['evidence_ids'])}")
            meta.paragraph_format.keep_with_next=True
            doc.add_paragraph(f['detail'])
            doc.add_paragraph('Proposed action: '+f['recommendation'])
    elif index in (3,4):
        for position,source in enumerate(DATA['domain_tables']):
            if source.get('report_section',4 if position==0 else 5)!=index+1:continue
            doc.add_heading(source['name'],level=2)
            if source.get('grain'):doc.add_paragraph('Record grain: '+source['grain'])
            if not source['rows']:
                doc.add_paragraph('No records supplied. Declared columns: '+', '.join(source['columns'])+'. An empty register does not establish a clean population.')
            for number,row in enumerate(source['rows'],1):
                label=doc.add_paragraph(f"{source['name']} / record {number} of {len(source['rows'])}")
                label.paragraph_format.keep_with_next=True
                # Long records may continue between field rows instead of overflowing a page.
                table(['Field','Value'], [(key,formatted(value,source.get('formats',{}).get(key))) for key,value in zip(source['columns'],row)],keep_together=False,widths=[2.0,5.0])
            missing = sum(value is None for row in source['rows'] for value in row)
            doc.add_paragraph(f"Domain record count: {len(source['rows'])}. Missing (null) cells: {missing}. Empty text follows the field-specific contract.")
    elif index == 5:
        table(['Evidence ID','Source or fixture','Exact locator','Claim class'], [[e['id'],e['title'],e['locator'],e['claim_class']] for e in DATA['evidence']],widths=[1.0,2.0,2.8,1.2])
    elif index == 6:
        for a in DATA['actions']:
            doc.add_paragraph(f"{a['id']} / {a['finding_id']}: {a['proposal']} Owner: {text(a['owner'])}. Due: {text(a['due_date'])}. State: {a['status']}")
    elif index == 7:
        for value in DATA['limitations']:
            doc.add_paragraph(value, style='List Bullet')
        table(['Check ID','Check','Result','Evidence'], [[q['id'],q['check'],q['result'],q['evidence']] for q in DATA['qa']])
footer = section.footer.paragraphs[0]
footer.add_run('Draft for human review | Page ')
field=OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
footer._p.append(field)
doc.save(OUTPUT)
print(str(OUTPUT))
