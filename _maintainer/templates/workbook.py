from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.properties import CalcProperties
from openpyxl.comments import Comment

DATA = None  # __DATA__
OUTPUT = Path(DATA['suite_id'] + '-review.xlsx')
wb = Workbook()
wb.remove(wb.active)
wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
wb.properties.creator = 'Markdown Analyst Toolkit'
wb.properties.title = DATA['title']
wb.properties.description = DATA['notice']
ink, muted, header, paper, stripe = '243641', '4E606A', '344E5F', 'F1F2EF', 'E2E8E9'
line = Side(style='thin', color='B8C3C8')

def literal(cell, value):
    if value is None:
        cell.value = None
        cell.comment = Comment('Missing value. Preserve the associated limitation or QA explanation; this is not zero.', 'Toolkit')
        return
    elif isinstance(value, (list, dict)):
        value = '; '.join(map(str, value)) if isinstance(value, list) else str(value)
    cell.value = value
    # Evidence text is data even when it starts with a spreadsheet formula marker.
    if isinstance(value, str):
        cell.data_type = 's'

def sheet(name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'B2'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = '1:1'
    ws.oddHeader.center.text = DATA['title'].replace('&', '&&')
    ws.oddFooter.left.text = 'Draft for human review'
    ws.oddFooter.right.text = 'Page &P of &N'
    return ws

def populate(ws, columns, rows, table_name=None, formats=None):
    formats = formats or {}
    for col, value in enumerate(columns, 1):
        literal(ws.cell(1, col), value)
    for row_number, values in enumerate(rows, 2):
        if len(values) != len(columns):
            raise ValueError(f'{ws.title}: row {row_number} has wrong column count')
        for col, value in enumerate(values, 1):
            literal(ws.cell(row_number, col), value)
    if table_name and rows:
        tab = Table(displayName=table_name, ref=ws.dimensions)
        tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(tab)
    ws.auto_filter.ref = ws.dimensions
    for col, name in enumerate(columns, 1):
        letter = ws.cell(1, col).column_letter
        long = any(word in name for word in ('detail', 'recommendation', 'proposal', 'method', 'locator', 'rationale'))
        ws.column_dimensions[letter].width = 44 if long else 30 if name in ('title','check') else max(14, min(25, len(name) + 3))
        for cells in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell = cells[0]
            if isinstance(cell.value, (float, int)):
                cell.number_format = formats.get(name) or ('$' + '#,##0.00;[Red]($#,##0.00);$0.00' if name.endswith('_usd') else '#,##0.00' if isinstance(cell.value, float) else '#,##0')
    return ws

summary = sheet('Summary')
summary_rows = [
    ['suite', DATA['title'], DATA['suite_id']],
    ['version', DATA['version'], 'Template version'],
    ['as_of', DATA['as_of'], 'Source cutoff, not file creation time'],
    ['status', 'DRAFT', 'Human disposition is not recorded'],
    ['scope', DATA['scope'], 'Coverage must reconcile before release'],
    ['decision_supported', DATA['decision'], 'Decision question; not approval'],
    ['data_label', DATA['notice'], 'Retain synthetic label for examples'],
]
summary_rows += [[m['label'], m['value'], m['method']] for m in DATA['metrics']]
summary_rows += [['limitation', value, 'Preserve in every exported format'] for value in DATA['limitations']]
populate(summary, ['measure', 'value', 'method_or_limit'], summary_rows)
for row_number, metric in enumerate(DATA['metrics'], 9):
    if isinstance(metric['value'], (int, float)):
        summary.cell(row_number, 2).number_format = metric.get('format', '#,##0.00')
summary.column_dimensions['A'].width = 30
summary.column_dimensions['B'].width = 70
summary.column_dimensions['C'].width = 80
dashboard = sheet('Dashboard')
dashboard.append(['severity', 'finding_count', 'denominator'])
finding_columns = ['id', 'title', 'severity', 'confidence', 'claim_class', 'evidence_ids', 'detail', 'recommendation']
findings = populate(sheet('Findings'), finding_columns,
                    [[f[k] for k in finding_columns] for f in DATA['findings']], 'FindingRegister')
for row, level in enumerate(['CRITICAL', 'HIGH', 'MEDIUM', 'LOW'], 2):
    literal(dashboard.cell(row, 1), level)
    dashboard.cell(row, 2, f'=COUNTIF(Findings!C2:C{max(2,len(DATA["findings"])+1)},A{row})')
    literal(dashboard.cell(row, 3), 'All supplied findings; not all underlying activity')
dashboard['A7'] = 'This is a generated snapshot. Regenerate after adding or removing rows.'
dashboard['A8'] = 'Open in a calculating spreadsheet application to refresh formula values.'
dashboard.merge_cells('A7:C7')
dashboard.merge_cells('A8:C8')
dashboard.column_dimensions['A'].width = 27
dashboard.column_dimensions['B'].width = 22
dashboard.column_dimensions['C'].width = 60
chart = BarChart()
chart.title = 'Findings by severity'
chart.y_axis.title = 'Findings'
chart.x_axis.title = 'Severity'
chart.add_data(Reference(dashboard, min_col=2, min_row=1, max_row=5), titles_from_data=True)
chart.set_categories(Reference(dashboard, min_col=1, min_row=2, max_row=5))
chart.style = 13
chart.y_axis.majorUnit = 1
chart.y_axis.scaling.min = 0
chart.series[0].graphicalProperties.solidFill = header
chart.series[0].graphicalProperties.line.solidFill = header
chart.width, chart.height = 23, 10
chart.legend = None
dashboard.add_chart(chart, 'A10')
ev_cols = ['id', 'title', 'locator', 'claim_class']
populate(sheet('Evidence'), ev_cols, [[e[k] for k in ev_cols] for e in DATA['evidence']], 'EvidenceRegister')
action_cols = ['action_id', 'finding_id', 'proposal', 'owner', 'due_date', 'status']
actions = populate(sheet('Actions'), action_cols,
                   [[a['id'], a['finding_id'], a['proposal'], a['owner'] or 'Not provided', a['due_date'] or 'Not provided', a['status']] for a in DATA['actions']], 'ActionRegister')
dv = DataValidation(type='list', formula1='"Proposed,Accepted,In progress,Blocked,Closed,Risk accepted"', allow_blank=False)
dv.errorTitle = 'Use a defined action state'
dv.error = 'Use the list. Closed requires closure evidence in the case record.'
dv.showErrorMessage = True
dv.errorStyle = 'stop'
actions.add_data_validation(dv)
dv.add(f'F2:F{max(2,len(DATA["actions"])+1)}')
qa_cols = ['check_id', 'check', 'result', 'evidence']
populate(sheet('QA'), qa_cols, [[q['id'], q['check'], q['result'], q['evidence']] for q in DATA['qa']], 'QARegister')
for i, t in enumerate(DATA['domain_tables'], 1):
    ws = populate(sheet(t['name']), t['columns'], t['rows'], f'DomainRegister{i}', t.get('formats'))
    stats_row = len(t['rows']) + 3
    literal(ws.cell(stats_row, 1), 'Domain record count')
    literal(ws.cell(stats_row, 2), len(t['rows']))
    literal(ws.cell(stats_row + 1, 1), 'Missing (null) cells')
    literal(ws.cell(stats_row + 1, 2), sum(value is None for row in t['rows'] for value in row))
for ws in wb:
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name='Arial', size=11, color=ink)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.fill = PatternFill('solid', fgColor=paper if cell.row % 2 else stripe)
            cell.border = Border(bottom=line)
            if cell.row == 1:
                cell.font = Font(name='Arial', size=11, color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor=header)
    ws.row_dimensions[1].height = 32
    if sum(dim.width for dim in ws.column_dimensions.values()) > 145:
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.print_title_cols = 'A:A'
    # No fixed body height: Excel can expand wrapped evidence without clipping.
    ws.print_area = ws.dimensions if ws.title != 'Dashboard' else 'A1:H31'
findings.conditional_formatting.add(f'C2:C{max(2,len(DATA["findings"])+1)}',
    CellIsRule(operator='equal', formula=['"HIGH"'], fill=PatternFill('solid',fgColor='F2DADB'),font=Font(color='7F3038')))
for ws in [summary, dashboard]:
    ws.freeze_panes = 'A2'
summary['A1'].comment = Comment(DATA['notice'] + ' Use file 05 for domain field meanings and file 09 for release checks.', 'Toolkit')
wb.save(OUTPUT)
print(str(OUTPUT))
