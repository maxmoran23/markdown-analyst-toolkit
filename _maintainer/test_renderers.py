"""Exercise real emitted files, including negative controls, not mocked libraries."""
import copy
import json
from pathlib import Path
import re
import runpy
import tempfile
import os
import unittest
from openpyxl import load_workbook
from docx import Document
from build import ROOT, COMMON, prepare, python_template, dashboard
from validate import validate_catalog

def run_python(source, directory, filename):
    path=directory/filename
    path.write_text(source)
    old=Path.cwd()
    try:
        os.chdir(directory)
        runpy.run_path(str(path))
    finally:
        os.chdir(old)

class RendererTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalogs=[json.loads(p.read_text()) for p in sorted((ROOT/'_maintainer/catalog').glob('*.json'))]
        if not cls.catalogs:
            raise RuntimeError('No catalogs: refusing vacuous success')

    def test_every_workbook_and_report_round_trip(self):
        for raw in self.catalogs:
            with self.subTest(suite=raw['slug']),tempfile.TemporaryDirectory() as tmp:
                data=prepare(raw);folder=Path(tmp)
                embedded=(ROOT/'suites'/data['slug']/'07-workbook-template.md').read_text()
                run_python(re.search(r'^```python\n(.*?)^```',embedded,re.M|re.S).group(1),folder,'workbook.py')
                wb=load_workbook(folder/(data['slug']+'-review.xlsx'),data_only=False)
                schemas={**COMMON,**{t['name']:t['columns'] for t in data['domain_tables']}}
                self.assertEqual(wb.sheetnames,list(schemas))
                for name,columns in schemas.items():
                    actual=[cell.value for cell in wb[name][1]][:len(columns)]
                    self.assertEqual(actual,columns)
                for t in data['domain_tables']:
                    rows=list(wb[t['name']].iter_rows(min_row=2,max_row=len(t['rows'])+1,values_only=True))
                    self.assertEqual(len(rows),len(t['rows']))
                    self.assertEqual(wb[t['name']].cell(len(t['rows'])+3,2).value,len(t['rows']))
                    self.assertEqual(wb[t['name']].cell(len(t['rows'])+4,2).value,sum(v is None for row in t['rows'] for v in row))
                    for emitted,original in zip(rows,t['rows']):
                        for actual, expected in zip(emitted, original):
                            if isinstance(expected, float):
                                self.assertAlmostEqual(actual, expected, places=12)
                            else:
                                self.assertEqual(actual, None if expected == "" else expected)
                    for column, fmt in t.get('formats', {}).items():
                        column_index = t['columns'].index(column) + 1
                        for cells in wb[t['name']].iter_rows(min_row=2,max_row=len(t['rows'])+1,min_col=column_index,max_col=column_index):
                            if isinstance(cells[0].value,(int,float)):
                                self.assertEqual(cells[0].number_format,fmt)
                self.assertEqual(wb['Findings'].max_row,len(data['findings'])+1)
                self.assertEqual(len(wb['Dashboard']._charts),1)
                self.assertTrue(wb['Dashboard']['B2'].value.startswith('=COUNTIF('))
                for ws in wb:
                    self.assertTrue(ws.freeze_panes)
                    self.assertTrue(ws.print_title_rows)
                embedded=(ROOT/'suites'/data['slug']/'06-report-template.md').read_text()
                run_python(re.search(r'^```python\n(.*?)^```',embedded,re.M|re.S).group(1),folder,'report.py')
                doc=Document(folder/(data['slug']+'-report.docx'))
                headings=[p.text for p in doc.paragraphs if p.style.name=='Heading 1']
                self.assertEqual(headings,data['report_sections'])
                all_text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                for f in data['findings']:
                    self.assertIn(f['id'],all_text)
                self.assertIn(data['notice'],all_text)
                evidence_tables=[t for t in doc.tables if [c.text for c in t.rows[0].cells]==['Evidence ID','Source or fixture','Exact locator','Claim class']]
                self.assertEqual(len(evidence_tables),1)
                self.assertEqual([[c.text for c in r.cells] for r in evidence_tables[0].rows[1:]],[[e['id'],e['title'],e['locator'],e['claim_class']] for e in data['evidence']])

    def test_report_preserves_declared_fx_precision(self):
        data=prepare(json.loads((ROOT/'_maintainer/catalog/source-wealth-funds.json').read_text()))
        table=data['domain_tables'][0];index=table['columns'].index('fx_usd_per_unit')
        table['rows'][0][index]=1.2345
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp);run_python(python_template('report.py',data),folder,'report.py')
            doc=Document(folder/(data['slug']+'-report.docx'))
            values=[r.cells[1].text for t in doc.tables for r in t.rows if len(r.cells)==2 and r.cells[0].text=='fx_usd_per_unit']
            self.assertEqual(values[0],'1.2345')
            self.assertIn('1.2000',values)

    def test_injected_evidence_stays_literal_in_all_formats(self):
        data=prepare(copy.deepcopy(self.catalogs[0]))
        attack='</script><script>window.pwned=true</script> =HYPERLINK("https://example.invalid")'
        data['evidence'][0]['title']=attack
        data['findings'][0]['title']='=1+1'
        rendered=dashboard(data)
        self.assertNotIn(attack,rendered)
        payload=re.search(r'<script id="report-data" type="application/json">(.*?)</script>',rendered,re.S).group(1)
        self.assertEqual(json.loads(payload)['evidence'][0]['title'],attack)
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)
            run_python(python_template('workbook.py',data),folder,'workbook.py')
            wb=load_workbook(folder/(data['slug']+'-review.xlsx'))
            self.assertEqual(wb['Findings']['B2'].value,'=1+1')
            self.assertEqual(wb['Findings']['B2'].data_type,'s')
            self.assertEqual(wb['Evidence']['B2'].value,attack)

    def test_missing_evidence_duplicate_ids_and_wrong_width_are_rejected(self):
        data=copy.deepcopy(self.catalogs[0])
        data['findings'][0]['evidence_ids']=['DOES-NOT-EXIST']
        self.assertTrue(any('unresolved evidence' in e for e in validate_catalog(data)))
        data=copy.deepcopy(self.catalogs[0]);data['evidence'].append(data['evidence'][0])
        self.assertIn('duplicate evidence ID',validate_catalog(data))
        data=copy.deepcopy(self.catalogs[0]);data['domain_tables'][0]['rows'][0].append('extra')
        self.assertTrue(any('width mismatch' in e for e in validate_catalog(data)))

    def test_empty_findings_do_not_create_fake_results(self):
        data=prepare(copy.deepcopy(self.catalogs[0]));data['findings']=[];data['actions']=[]
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp)
            run_python(python_template('workbook.py',data),folder,'workbook.py')
            wb=load_workbook(folder/(data['slug']+'-review.xlsx'))
            self.assertEqual(wb['Findings'].max_row,1)
            self.assertEqual(wb['Actions'].max_row,1)
            self.assertNotIn('FindingRegister',wb['Findings'].tables)
        self.assertEqual(json.loads(re.search(r'type="application/json">(.*?)</script>',dashboard(data),re.S).group(1))['findings'],[])

if __name__=='__main__':
    unittest.main(verbosity=2)
